from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Callable

import numpy as np
from openpyxl import Workbook, load_workbook

from .measurements import (
    apply_spine_volume_filter,
    distribution_summary_rows,
    filtered_measurement_result,
    load_distribution_preview,
    spine_volume_distribution_rows,
    spine_volume_filter_settings,
)


Progress = Callable[[str, int, int, str], None]
BIN_COLORS = (
    (35, 0, 75), (75, 3, 110), (112, 14, 117), (147, 37, 103), (177, 63, 82),
    (204, 93, 58), (224, 127, 36), (239, 167, 25), (246, 210, 42), (240, 249, 33),
)


def _cell(value: object) -> object:
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def _columns(rows: list[dict[str, object]]) -> list[str]:
    columns: list[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    return columns


def _group_summary(specimen_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    output: list[dict[str, object]] = []
    groups = sorted({str(row.get("experimental_group", "")) for row in specimen_rows})
    excluded = {"experimental_group", "specimen_id", "average_protein_distribution"}
    numeric = sorted(
        {
            key
            for row in specimen_rows
            for key, value in row.items()
            if key not in excluded and isinstance(value, (int, float)) and not isinstance(value, bool)
        }
    )
    for group in groups:
        members = [row for row in specimen_rows if str(row.get("experimental_group", "")) == group]
        for metric in numeric:
            values = [float(row[metric]) for row in members if row.get(metric) is not None]
            if not values:
                continue
            sd = float(np.std(values, ddof=1)) if len(values) > 1 else 0.0
            output.append(
                {
                    "experimental_group": group,
                    "metric": metric,
                    "n_specimens": len(values),
                    "mean": float(np.mean(values)),
                    "sd": sd,
                    "sem": sd / np.sqrt(len(values)),
                }
            )
    return output


def collect_export_tables(manifest: dict[str, object]) -> dict[str, list[dict[str, object]]]:
    results: list[dict[str, object]] = []
    volume_audit: list[dict[str, object]] = []
    for index, specimen_record in enumerate(manifest["specimens"]):
        if specimen_record["checkpoints"].get("measurements", {}).get("state") != "complete":
            continue
        result, audit = filtered_measurement_result(manifest, index)
        distribution_by_spine = {
            int(row.get("spine_id") or 0): row
            for row in result.get("distribution_rows", [])
        }
        for row in audit:
            distribution = distribution_by_spine.get(int(row.get("spine_id") or 0), {})
            row.update(
                {
                    "distribution_included": distribution.get("distribution_included"),
                    "distribution_reviewed": distribution.get(
                        "distribution_reviewed", row.get("distribution_reviewed")
                    ),
                    "distribution_axis_status": distribution.get(
                        "distribution_axis_status"
                    ),
                    "distribution_review_required": distribution.get(
                        "distribution_review_required"
                    ),
                    "centerline_endpoint_source": distribution.get(
                        "centerline_endpoint_source"
                    ),
                    "review_note": distribution.get(
                        "review_note", row.get("validity_note", "")
                    ),
                }
            )
        results.append(result)
        volume_audit.extend(audit)
    return collect_export_tables_from_results(manifest, results, volume_audit=volume_audit)


def collect_export_tables_from_results(
    manifest: dict[str, object],
    results: list[dict[str, object]],
    *,
    volume_audit: list[dict[str, object]] | None = None,
) -> dict[str, list[dict[str, object]]]:
    filter_enabled, cutoff = spine_volume_filter_settings(manifest)
    specimen = [row for result in results for row in result.get("specimen_rows", [])]
    roi = [row for result in results for row in result.get("roi_rows", [])]
    dendrite = [row for result in results for row in result.get("dendrite_rows", [])]
    all_spines = [row for result in results for row in result.get("spine_rows", [])]
    spine = [row for row in all_spines if not bool(row.get("volume_filter_excluded", False))]
    clusters = [
        row
        for result in results
        for row in result.get("cluster_rows", [])
        if row.get("row_type") == "individual_cluster"
        and not bool(row.get("volume_filter_excluded", False))
    ]
    sums = [
        row
        for result in results
        for row in result.get("cluster_rows", [])
        if row.get("row_type") == "spine_cluster_sum"
        and not bool(row.get("volume_filter_excluded", False))
    ]
    distributions = [
        row
        for result in results
        for row in result.get("distribution_rows", [])
        if not bool(row.get("volume_filter_excluded", False))
    ]
    distribution_specimen, distribution_group = distribution_summary_rows(results)
    excluded = [
        row
        for row in distributions
        if bool(row.get("spine_valid", True)) and not bool(row.get("distribution_included", False))
    ]
    invalid_ids = {
        (str(row.get("experimental_group", "")), str(row.get("specimen_id", "")), int(row["spine_id"]))
        for row in spine
        if not bool(row.get("manual_spine_valid", row.get("spine_valid", True)))
    }
    invalid = [
        row
        for row in spine
        if (str(row.get("experimental_group", "")), str(row.get("specimen_id", "")), int(row["spine_id"])) in invalid_ids
    ]
    reviewed_spines = [
        row
        for row in spine
        if not bool(row.get("volume_filter_excluded", False))
        and (bool(row.get("validity_reviewed", False))
        or bool(row.get("distribution_reviewed", False))
        )
    ]
    settings = [
        {"section": "calibration", **dict(manifest["calibration"])},
        {"section": "measurements", **dict(manifest["measurements"]["settings"])},
        {
            "section": "distribution",
            "bin_count": 10,
            "orientation": "largest shaft contact to longest curved distal skeleton path",
            "aggregation": "spine means within specimen, then specimen means within group",
            "group_error_bars": "SEM across specimen means",
            "ratio_units": "fraction 0-1",
        },
        {
            "section": "spine_volume_filter",
            "enabled": filter_enabled,
            "cutoff_um3": cutoff,
            "rule": "exclude volume_um3 < cutoff_um3; equality retained",
            "force_keep_scope": "volume rule only; manually invalid spines remain invalid",
        },
    ]
    excluded_specimens = [
        {
            "experimental_group": specimen.get("experimental_group", ""),
            "specimen_id": specimen.get("specimen_id", ""),
            "reason": specimen.get("analysis", {}).get("exclusion_reason", ""),
        }
        for specimen in manifest["specimens"]
        if bool(specimen.get("analysis", {}).get("excluded", False))
    ]
    return {
        "Specimen_Master": sorted(specimen, key=lambda row: (str(row.get("experimental_group")), str(row.get("specimen_id")))),
        "ROI_Master": sorted(roi, key=lambda row: (str(row.get("experimental_group")), str(row.get("specimen_id")), int(row.get("roi_id", 0)))),
        "Dendrite_Master": sorted(dendrite, key=lambda row: (str(row.get("experimental_group")), str(row.get("specimen_id")), int(row.get("dendrite_id", 0)))),
        "Spine_Master": sorted(spine, key=lambda row: (str(row.get("experimental_group")), str(row.get("specimen_id")), int(row.get("spine_id", 0)))),
        "Cluster_Individual": clusters,
        "Cluster_Sums": sums,
        "Distribution_Individual": distributions,
        "Distribution_Specimen": distribution_specimen,
        "Distribution_Group": distribution_group,
        "Distribution_Excluded": excluded,
        "Invalid_Spines": invalid,
        "Spine_Review_Audit": reviewed_spines,
        "Volume_Filtered_Spines": list(volume_audit or []),
        "Excluded_Specimens": excluded_specimens,
        "Group_Summary": _group_summary(specimen),
        "Spine_Volume_Distribution": spine_volume_distribution_rows(
            results, filter_enabled=filter_enabled, cutoff_um3=cutoff
        ),
        "Settings": settings,
    }


def _write_workbook(path: Path, tables: dict[str, list[dict[str, object]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in tables.items():
        sheet = workbook.create_sheet(name[:31])
        columns = _columns(rows)
        if columns:
            sheet.append(columns)
            for row in rows:
                sheet.append([_cell(row.get(column)) for column in columns])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        else:
            sheet.append(["No rows"])
    workbook.save(path)


def _write_csvs(directory: Path, tables: dict[str, list[dict[str, object]]]) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for name, rows in tables.items():
        path = directory / f"{name}.csv"
        columns = _columns(rows)
        with path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
            if columns:
                writer.writeheader()
                for row in rows:
                    writer.writerow({key: _cell(row.get(key)) for key in columns})
        written.append(path)
    return written


def _worksheet_rows(workbook, name: str) -> list[dict[str, object]]:  # type: ignore[no-untyped-def]
    if name not in workbook.sheetnames:
        return []
    values = list(workbook[name].iter_rows(values_only=True))
    if not values or not values[0] or values[0][0] == "No rows":
        return []
    headers = [str(value) if value is not None else "" for value in values[0]]
    return [
        {header: value for header, value in zip(headers, row) if header}
        for row in values[1:]
        if any(value is not None for value in row)
    ]


def _worksheet_columns(workbook, name: str) -> set[str]:  # type: ignore[no-untyped-def]
    if name not in workbook.sheetnames:
        return set()
    first = next(workbook[name].iter_rows(min_row=1, max_row=1, values_only=True), ())
    if not first or first[0] == "No rows":
        return set()
    return {str(value) for value in first if value is not None and str(value)}


def inspect_exported_measurement_workbook(path: str | Path) -> dict[str, object]:
    """Validate a Synpo workbook and return rows needed by the filter dialog."""
    source = Path(path).resolve()
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        rows = _worksheet_rows(workbook, "Spine_Master")
        required = {
            "experimental_group",
            "specimen_id",
            "spine_id",
            "dendrite_id",
            "volume_um3",
            "spine_valid",
        }
        columns = _worksheet_columns(workbook, "Spine_Master")
        missing = sorted(required - columns)
        if missing:
            raise ValueError(
                "Spine_Master is missing required column(s): " + ", ".join(missing)
            )
        settings = _worksheet_rows(workbook, "Settings")
        volume_setting = next(
            (row for row in settings if row.get("section") == "spine_volume_filter"),
            {},
        )
        return {
            "path": str(source),
            "spines": rows,
            "cutoff_um3": float(volume_setting.get("cutoff_um3") or 0.0),
            "filter_enabled": bool(volume_setting.get("enabled", False)),
            "sheet_names": list(workbook.sheetnames),
            "sheet_columns": {
                name: sorted(_worksheet_columns(workbook, name))
                for name in workbook.sheetnames
            },
        }
    finally:
        workbook.close()


def _group_workbook_results(
    tables: dict[str, list[dict[str, object]]],
) -> list[dict[str, object]]:
    keys = sorted(
        {
            (str(row.get("experimental_group", "")), str(row.get("specimen_id", "")))
            for row in tables.get("Spine_Master", [])
        }
    )
    results: list[dict[str, object]] = []
    table_map = {
        "Specimen_Master": "specimen_rows",
        "ROI_Master": "roi_rows",
        "Dendrite_Master": "dendrite_rows",
        "Spine_Master": "spine_rows",
        "Distribution_Individual": "distribution_rows",
    }
    for group, specimen in keys:
        result: dict[str, object] = {}
        for table_name, result_name in table_map.items():
            result[result_name] = [
                dict(row)
                for row in tables.get(table_name, [])
                if str(row.get("experimental_group", "")) == group
                and str(row.get("specimen_id", "")) == specimen
            ]
        if not result["specimen_rows"]:
            result["specimen_rows"] = [
                {"experimental_group": group, "specimen_id": specimen}
            ]
        clusters: list[dict[str, object]] = []
        for table_name, row_type in (
            ("Cluster_Individual", "individual_cluster"),
            ("Cluster_Sums", "spine_cluster_sum"),
        ):
            for source in tables.get(table_name, []):
                if str(source.get("experimental_group", "")) != group or str(
                    source.get("specimen_id", "")
                ) != specimen:
                    continue
                row = dict(source)
                row.setdefault("row_type", row_type)
                clusters.append(row)
        result["cluster_rows"] = clusters
        results.append(result)
    return results


def _replace_sheet(workbook, name: str, rows: list[dict[str, object]]) -> None:  # type: ignore[no-untyped-def]
    index = workbook.sheetnames.index(name) if name in workbook.sheetnames else len(workbook.sheetnames)
    if name in workbook.sheetnames:
        workbook.remove(workbook[name])
    sheet = workbook.create_sheet(name[:31], index)
    columns = _columns(rows)
    if not columns:
        sheet.append(["No rows"])
        return
    sheet.append(columns)
    for row in rows:
        sheet.append([_cell(row.get(column)) for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _write_workbook_csv_directory(workbook, directory: Path) -> list[Path]:  # type: ignore[no-untyped-def]
    directory.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    used: set[str] = set()
    for sheet in workbook.worksheets:
        stem = re.sub(r"[^A-Za-z0-9_.-]+", "_", sheet.title).strip("._") or "Sheet"
        candidate = stem
        suffix = 2
        while candidate.lower() in used:
            candidate = f"{stem}_{suffix}"
            suffix += 1
        used.add(candidate.lower())
        path = directory / f"{candidate}.csv"
        with path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.writer(stream)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow([_cell(value) for value in row])
        written.append(path)
    return written


def filter_exported_measurement_workbook(
    source_path: str | Path,
    output_path: str | Path,
    *,
    cutoff_um3: float,
    force_keep_keys: set[tuple[str, str, int]] | None = None,
) -> dict[str, object]:
    """Filter and exactly recalculate supported tables in an exported workbook."""
    source = Path(source_path).resolve()
    output = Path(output_path).resolve()
    if source == output:
        raise ValueError("Choose a new output file; the source workbook is never overwritten.")
    inspection = inspect_exported_measurement_workbook(source)
    if bool(inspection.get("filter_enabled", False)):
        raise ValueError(
            "This workbook is already volume-filtered and no longer contains the omitted "
            "detail rows. Select the original unfiltered workbook so changing the cutoff "
            "remains reversible."
        )
    force_keep = force_keep_keys or set()
    values = load_workbook(source, read_only=True, data_only=True)
    try:
        tables = {name: _worksheet_rows(values, name) for name in values.sheetnames}
    finally:
        values.close()
    sheet_columns = {
        str(name): set(columns)
        for name, columns in dict(inspection["sheet_columns"]).items()
    }
    table_requirements = {
        "Spine_Master": {"experimental_group", "specimen_id", "dendrite_id", "spine_id", "volume_um3", "spine_valid"},
        "Cluster_Individual": {"experimental_group", "specimen_id", "spine_id"},
        "Cluster_Sums": {"experimental_group", "specimen_id", "spine_id"},
        "Distribution_Individual": {"experimental_group", "specimen_id", "spine_id", "distribution_included"},
        "Dendrite_Master": {"experimental_group", "specimen_id", "roi_id", "dendrite_id", "length_um"},
        "ROI_Master": {"experimental_group", "specimen_id", "roi_id"},
        "Specimen_Master": {"experimental_group", "specimen_id"},
    }
    table_issues: dict[str, str] = {}
    for name, required in table_requirements.items():
        if name not in inspection["sheet_names"]:
            table_issues[name] = f"missing detailed sheet {name}"
            continue
        missing_columns = sorted(required - sheet_columns.get(name, set()))
        if missing_columns:
            table_issues[name] = "missing column(s): " + ", ".join(missing_columns)
            if name != "Spine_Master":
                tables[name] = []
    spine_summary_columns = {
        "roi_id",
        "has_protein_cluster",
        "cluster_to_spine_volume_ratio",
    }
    cluster_summary_columns = {"dendrite_id", "roi_id", "volume_inside_spine_um3"}
    distribution_summary_columns = {
        "dendrite_id",
        "roi_id",
        "distribution_included",
        *{f"bin_{index:02d}_ratio" for index in range(1, 11)},
    }
    summary_missing: list[str] = []
    for name in ("Dendrite_Master", "Cluster_Individual", "Distribution_Individual"):
        if name in table_issues:
            summary_missing.append(table_issues[name])
    for name, required in (
        ("Spine_Master", spine_summary_columns),
        ("Cluster_Individual", cluster_summary_columns),
        ("Distribution_Individual", distribution_summary_columns),
    ):
        missing = sorted(required - sheet_columns.get(name, set()))
        if missing:
            summary_missing.append(f"{name} missing column(s): {', '.join(missing)}")
    summary_ready = not summary_missing
    distribution_ready = (
        "Distribution_Individual" not in table_issues
        and not (
            {f"bin_{index:02d}_ratio" for index in range(1, 11)}
            - sheet_columns.get("Distribution_Individual", set())
        )
    )
    review_audit_columns = {"validity_reviewed"}
    review_audit_ready = not (
        review_audit_columns - sheet_columns.get("Spine_Master", set())
    )
    audit_columns = {
        "roi_id",
        "has_protein_cluster",
        "included_cluster_count",
        "inside_cluster_volume_sum_um3",
    }
    missing_audit_columns = sorted(
        audit_columns - sheet_columns.get("Spine_Master", set())
    )
    raw_results = _group_workbook_results(tables)
    filtered_results: list[dict[str, object]] = []
    audit: list[dict[str, object]] = []
    for result in raw_results:
        identity = result.get("specimen_rows", [{}])[0]
        group = str(identity.get("experimental_group", ""))
        specimen = str(identity.get("specimen_id", ""))
        ids = {
            spine_id
            for key_group, key_specimen, spine_id in force_keep
            if key_group == group and key_specimen == specimen
        }
        filtered, excluded = apply_spine_volume_filter(
            result,
            enabled=True,
            cutoff_um3=cutoff_um3,
            force_keep_ids=ids,
            refresh_summaries=summary_ready,
        )
        distribution_by_spine = {
            int(row.get("spine_id") or 0): row
            for row in filtered.get("distribution_rows", [])
        }
        for row in excluded:
            distribution = distribution_by_spine.get(int(row.get("spine_id") or 0), {})
            row.update(
                {
                    "distribution_included": distribution.get("distribution_included"),
                    "distribution_reviewed": distribution.get("distribution_reviewed"),
                    "distribution_axis_status": distribution.get("distribution_axis_status"),
                    "distribution_review_required": distribution.get("distribution_review_required"),
                    "centerline_endpoint_source": distribution.get("centerline_endpoint_source"),
                    "review_note": distribution.get("review_note", row.get("validity_note", "")),
                }
            )
        filtered_results.append(filtered)
        audit.extend(excluded)

    manifest = {
        "measurements": {
            "settings": {
                "spine_volume_filter_enabled": True,
                "spine_volume_filter_cutoff_um3": float(cutoff_um3),
            }
        },
        "calibration": {},
        "specimens": [],
    }
    recalculated = collect_export_tables_from_results(
        manifest, filtered_results, volume_audit=audit
    )
    available = set(inspection["sheet_names"]) - set(table_issues)
    if summary_ready:
        available.add("__summary_inputs__")
    if distribution_ready:
        available.add("__distribution_inputs__")
    if review_audit_ready:
        available.add("__review_audit_inputs__")
    dependencies = {
        "Spine_Master": {"Spine_Master"},
        "Cluster_Individual": {"Spine_Master", "Cluster_Individual"},
        "Cluster_Sums": {"Spine_Master", "Cluster_Sums"},
        "Distribution_Individual": {"Spine_Master", "Distribution_Individual"},
        "Distribution_Specimen": {"Spine_Master", "Distribution_Individual", "__distribution_inputs__"},
        "Distribution_Group": {"Spine_Master", "Distribution_Individual", "__distribution_inputs__"},
        "Distribution_Excluded": {"Spine_Master", "Distribution_Individual"},
        "Dendrite_Master": {"Spine_Master", "Dendrite_Master", "Cluster_Individual", "Distribution_Individual", "__summary_inputs__"},
        "ROI_Master": {"Spine_Master", "ROI_Master", "Dendrite_Master", "Cluster_Individual", "Distribution_Individual", "__summary_inputs__"},
        "Specimen_Master": {"Spine_Master", "Specimen_Master", "Dendrite_Master", "Cluster_Individual", "Distribution_Individual", "__summary_inputs__"},
        "Group_Summary": {"Spine_Master", "Specimen_Master", "Dendrite_Master", "Cluster_Individual", "Distribution_Individual", "__summary_inputs__"},
        "Invalid_Spines": {"Spine_Master"},
        "Spine_Review_Audit": {"Spine_Master", "__review_audit_inputs__"},
        "Volume_Filtered_Spines": {"Spine_Master"},
        "Spine_Volume_Distribution": {"Spine_Master"},
    }
    report: list[dict[str, object]] = []
    workbook = load_workbook(source, data_only=False)
    for name, needed in dependencies.items():
        missing = sorted(needed - available)
        if missing:
            status = "unavailable"
            reasons = [
                "; ".join(summary_missing)
                if item == "__summary_inputs__"
                else "Distribution_Individual is missing one or more bin_01_ratio through bin_10_ratio columns"
                if item == "__distribution_inputs__"
                else "Spine_Master is missing validity_reviewed"
                if item == "__review_audit_inputs__"
                else table_issues.get(item, f"missing detailed sheet {item}")
                for item in missing
            ]
            detail = "Exact recalculation unavailable: " + "; ".join(
                dict.fromkeys(reasons)
            )
            if name in workbook.sheetnames:
                _replace_sheet(
                    workbook,
                    name,
                    [{"status": "NOT RECALCULATED", "reason": detail}],
                )
        else:
            status = "recalculated"
            detail = "Filtered and recalculated exactly from detailed source sheets."
            _replace_sheet(workbook, name, recalculated.get(name, []))
            if name == "Volume_Filtered_Spines" and missing_audit_columns:
                status = "recalculated_with_unavailable_fields"
                detail = (
                    "Exclusion rows were recalculated exactly, but legacy Spine_Master "
                    "does not provide: " + ", ".join(missing_audit_columns)
                )
        report.append({"table": name, "status": status, "detail": detail})

    report.append(
        {
            "table": "Settings",
            "status": "updated",
            "detail": "Original settings retained and the applied volume-filter rule appended.",
        }
    )
    if "Excluded_Specimens" in workbook.sheetnames:
        report.append(
            {
                "table": "Excluded_Specimens",
                "status": "preserved",
                "detail": "Not spine-volume dependent; preserved unchanged.",
            }
        )
    known = set(dependencies) | {"Settings", "Excluded_Specimens", "Compatibility_Report"}
    for name in workbook.sheetnames:
        if name not in known:
            report.append(
                {
                    "table": name,
                    "status": "preserved_unrecognized",
                    "detail": "User-added or unrecognized sheet preserved unchanged.",
                }
            )
    report.append(
        {
            "table": "PDF exports",
            "status": "unavailable",
            "detail": "Standalone workbooks contain no source pixels/cache; PDFs were not regenerated.",
        }
    )

    original_settings = [
        row for row in tables.get("Settings", []) if row.get("section") != "spine_volume_filter"
    ]
    original_settings.append(
        {
            "section": "spine_volume_filter",
            "enabled": True,
            "cutoff_um3": float(cutoff_um3),
            "rule": "exclude volume_um3 < cutoff_um3; equality retained",
            "source_workbook": str(source),
        }
    )
    _replace_sheet(workbook, "Settings", original_settings)
    _replace_sheet(workbook, "Compatibility_Report", report)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    csv_directory = output.parent / f"{output.stem}_csv"
    csv_paths = _write_workbook_csv_directory(workbook, csv_directory)
    workbook.close()
    load_workbook(output, read_only=True).close()
    return {
        "workbook": str(output),
        "csv_directory": str(csv_directory),
        "source_workbook": str(source),
        "cutoff_um3": float(cutoff_um3),
        "excluded_spine_count": len(audit),
        "force_kept_count": len(force_keep),
        "compatibility_report": report,
        "csv_count": len(csv_paths),
        "pdfs_regenerated": False,
        "verified": True,
    }


def _pdf_pages_matplotlib_legacy(
    manifest: dict[str, object],
    path: Path,
    selected: list[tuple[int, dict[str, object]]],
    group_rows: list[dict[str, object]],
    margin_um: float,
    *,
    include_group_summary: bool,
    progress: Progress | None,
) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    colors = np.asarray(BIN_COLORS, dtype=np.float64) / 255.0
    with PdfPages(path) as pdf:
        wrote_page = False
        if include_group_summary:
            for group in group_rows:
                means = np.asarray(
                    [
                        np.nan if group.get(f"bin_{index:02d}_mean") is None else float(group[f"bin_{index:02d}_mean"])
                        for index in range(1, 11)
                    ],
                    dtype=np.float64,
                )
                sem = np.asarray(
                    [
                        np.nan if group.get(f"bin_{index:02d}_sem") is None else float(group[f"bin_{index:02d}_sem"])
                        for index in range(1, 11)
                    ],
                    dtype=np.float64,
                )
                figure, axis = plt.subplots(figsize=(9, 6))
                x_values = np.arange(1, 11, dtype=np.float64)
                axis.plot(x_values, means, marker="o")
                for x_value, mean, error in zip(x_values, means, sem):
                    if not np.isfinite(mean) or not np.isfinite(error):
                        continue
                    axis.plot([x_value, x_value], [mean - error, mean + error], color="C0")
                    axis.plot([x_value - 0.08, x_value + 0.08], [mean - error, mean - error], color="C0")
                    axis.plot([x_value - 0.08, x_value + 0.08], [mean + error, mean + error], color="C0")
                axis.set(xlabel="Spine part (shaft → tip)", ylabel="Cluster volume / spine-part volume", title=f"{group['experimental_group']} — specimen-weighted mean ± SEM")
                axis.set_xlim(0.5, 10.5)
                axis.set_ylim(bottom=0)
                axis.grid(alpha=0.25)
                pdf.savefig(figure, bbox_inches="tight")
                plt.close(figure)
                wrote_page = True
        for position, (specimen_index, row) in enumerate(selected):
            preview = load_distribution_preview(
                manifest, specimen_index, int(row["spine_id"]), margin_um=margin_um
            )
            figure = plt.figure(figsize=(11.7, 8.3))
            grid = figure.add_gridspec(2, 2, height_ratios=(3, 2))
            dendrite_axis = figure.add_subplot(grid[0, 0])
            protein_axis = figure.add_subplot(grid[0, 1])
            profile_axis = figure.add_subplot(grid[1, :])
            for axis, raw, bins, title in (
                (dendrite_axis, preview.dendrite_projection, preview.spine_bins_projection, "Dendrite/spine channel"),
                (protein_axis, preview.protein_projection, preview.cluster_bins_projection, "Protein channel"),
            ):
                axis.imshow(raw, cmap="gray", vmin=np.percentile(raw, 0.5), vmax=np.percentile(raw, 99.8))
                overlay = np.zeros((*bins.shape, 4), dtype=np.float64)
                for bin_index in range(1, 11):
                    overlay[bins == bin_index] = (*colors[bin_index - 1, :3], 0.55 if axis is dendrite_axis else 0.85)
                axis.imshow(overlay)
                if np.any(bins):
                    axis.contour(bins > 0, levels=[0.5], colors="white", linewidths=0.7)
                axis.set_title(title)
                axis.axis("off")
            if preview.axis_xy:
                dendrite_axis.plot(
                    [point[0] for point in preview.axis_xy],
                    [point[1] for point in preview.axis_xy],
                    color="white",
                    linewidth=1.2,
                )
            if preview.base_point_local_zyx is not None:
                dendrite_axis.scatter(
                    [preview.base_point_local_zyx[2]],
                    [preview.base_point_local_zyx[1]],
                    s=42,
                    c="#20d060",
                    edgecolors="black",
                    linewidths=0.6,
                    label="automatic base",
                    zorder=5,
                )
            if preview.endpoint_local_zyx is not None:
                dendrite_axis.scatter(
                    [preview.endpoint_local_zyx[2]],
                    [preview.endpoint_local_zyx[1]],
                    s=42,
                    c="#ed32c8",
                    edgecolors="black",
                    linewidths=0.6,
                    label=f"{row.get('centerline_endpoint_source', 'automatic')} endpoint",
                    zorder=5,
                )
            if preview.base_point_local_zyx is not None or preview.endpoint_local_zyx is not None:
                dendrite_axis.legend(loc="lower right", fontsize=7)
            ratios = [row.get(f"bin_{index:02d}_ratio") for index in range(1, 11)]
            profile_axis.plot(range(1, 11), ratios, marker="o")
            profile_axis.set(xlabel="Spine part (shaft → tip)", ylabel="Cluster volume / spine-part volume")
            profile_axis.set_ylim(bottom=0)
            profile_axis.grid(alpha=0.25)
            figure.suptitle(
                f"{row['experimental_group']} | {row['specimen_id']} | spine {row['spine_id']}\n"
                f"Axis: {row['distribution_axis_status']} | included: {row.get('distribution_included')} | "
                f"valid: {row.get('spine_valid')} | endpoint: "
                f"{row.get('centerline_endpoint_source', 'automatic')} | {row.get('review_note', '')}"
            )
            pdf.savefig(figure, bbox_inches="tight")
            plt.close(figure)
            wrote_page = True
            if progress:
                progress("Exporting validation PDF", position + 1, len(selected), f"Spine {row['spine_id']}")
        if not wrote_page:
            figure, axis = plt.subplots(figsize=(9, 6))
            axis.axis("off")
            axis.text(
                0.5,
                0.5,
                "No spines matched this optional PDF category.",
                ha="center",
                va="center",
                fontsize=14,
            )
            pdf.savefig(figure, bbox_inches="tight")
            plt.close(figure)


def _pdf_pages(
    manifest: dict[str, object],
    path: Path,
    selected: list[tuple[int, dict[str, object]]],
    group_rows: list[dict[str, object]],
    margin_um: float,
    *,
    include_group_summary: bool,
    progress: Progress | None,
) -> None:
    """Render validation pages without depending on matplotlib native DLLs."""
    from PIL import Image, ImageDraw, ImageFont

    page_size = (1400, 1000)
    font = ImageFont.load_default()

    def page_and_draw():  # type: ignore[no-untyped-def]
        page = Image.new("RGB", page_size, "white")
        return page, ImageDraw.Draw(page)

    def safe_text(value: object) -> str:
        return str(value).encode("ascii", "replace").decode("ascii")

    def draw_chart(draw, values: np.ndarray, errors: np.ndarray | None, box):  # type: ignore[no-untyped-def]
        left, top, right, bottom = box
        finite = values[np.isfinite(values)]
        maximum = float(np.max(finite)) if finite.size else 1.0
        if errors is not None:
            upper = values + np.nan_to_num(errors, nan=0.0)
            upper = upper[np.isfinite(upper)]
            maximum = max(maximum, float(np.max(upper)) if upper.size else 1.0)
        maximum = max(maximum, 1e-12)
        draw.rectangle(box, outline=(80, 80, 80), width=2)
        points: list[tuple[int, int]] = []
        for index, value in enumerate(values):
            if not np.isfinite(value):
                continue
            x = int(left + (index + 0.5) * (right - left) / len(values))
            y = int(bottom - float(value) / maximum * (bottom - top))
            points.append((x, y))
            if errors is not None and np.isfinite(errors[index]):
                delta = int(float(errors[index]) / maximum * (bottom - top))
                draw.line((x, y - delta, x, y + delta), fill=(35, 100, 175), width=2)
                draw.line((x - 5, y - delta, x + 5, y - delta), fill=(35, 100, 175), width=2)
                draw.line((x - 5, y + delta, x + 5, y + delta), fill=(35, 100, 175), width=2)
            draw.ellipse((x - 4, y - 4, x + 4, y + 4), fill=(35, 100, 175))
            draw.text((x - 3, bottom + 8), str(index + 1), fill="black", font=font)
        if len(points) > 1:
            draw.line(points, fill=(35, 100, 175), width=3)

    def projection(raw: np.ndarray, bins: np.ndarray, alpha: float):  # type: ignore[no-untyped-def]
        array = np.asarray(raw, dtype=np.float64)
        low, high = np.percentile(array, (0.5, 99.8))
        gray = np.clip((array - low) / max(float(high - low), 1e-12) * 255.0, 0, 255)
        rgb = np.repeat(gray[..., None], 3, axis=2)
        for bin_index, color in enumerate(BIN_COLORS, start=1):
            mask = np.asarray(bins) == bin_index
            rgb[mask] = (1.0 - alpha) * rgb[mask] + alpha * np.asarray(color)
        return Image.fromarray(np.clip(rgb, 0, 255).astype(np.uint8), mode="RGB")

    def paste_projection(page, source, box):  # type: ignore[no-untyped-def]
        left, top, right, bottom = box
        available = (right - left, bottom - top)
        resized = source.copy()
        resized.thumbnail(available, Image.Resampling.NEAREST)
        x = left + (available[0] - resized.width) // 2
        y = top + (available[1] - resized.height) // 2
        page.paste(resized, (x, y))
        return x, y, resized.width / source.width, resized.height / source.height

    pages = []
    if include_group_summary:
        for group in group_rows:
            means = np.asarray([
                np.nan if group.get(f"bin_{index:02d}_mean") is None else float(group[f"bin_{index:02d}_mean"])
                for index in range(1, 11)
            ])
            sem = np.asarray([
                np.nan if group.get(f"bin_{index:02d}_sem") is None else float(group[f"bin_{index:02d}_sem"])
                for index in range(1, 11)
            ])
            page, draw = page_and_draw()
            title = f"{group['experimental_group']} - specimen-weighted mean +/- SEM"
            draw.text((70, 45), safe_text(title), fill="black", font=font)
            draw_chart(draw, means, sem, (100, 130, 1320, 850))
            draw.text((520, 900), "Spine part (shaft -> tip)", fill="black", font=font)
            pages.append(page)

    for position, (specimen_index, row) in enumerate(selected):
        preview = load_distribution_preview(
            manifest, specimen_index, int(row["spine_id"]), margin_um=margin_um
        )
        page, draw = page_and_draw()
        title = (
            f"{row['experimental_group']} | {row['specimen_id']} | spine {row['spine_id']} | "
            f"axis: {row['distribution_axis_status']} | included: {row.get('distribution_included')} | "
            f"valid: {row.get('spine_valid')} | endpoint: {row.get('centerline_endpoint_source', 'automatic')}"
        )
        draw.text((45, 30), safe_text(title), fill="black", font=font)
        draw.text((45, 55), safe_text(row.get("review_note", "")), fill="black", font=font)
        left_box = (40, 105, 680, 570)
        right_box = (720, 105, 1360, 570)
        left_image = projection(preview.dendrite_projection, preview.spine_bins_projection, 0.55)
        right_image = projection(preview.protein_projection, preview.cluster_bins_projection, 0.85)
        left_x, left_y, scale_x, scale_y = paste_projection(page, left_image, left_box)
        paste_projection(page, right_image, right_box)
        draw.text((40, 85), "Dendrite/spine channel", fill="black", font=font)
        draw.text((720, 85), "Protein channel", fill="black", font=font)
        if preview.axis_xy:
            axis_points = [
                (left_x + int(point[0] * scale_x), left_y + int(point[1] * scale_y))
                for point in preview.axis_xy
            ]
            if len(axis_points) > 1:
                draw.line(axis_points, fill="white", width=3)
        for point, color in (
            (preview.base_point_local_zyx, (32, 208, 96)),
            (preview.endpoint_local_zyx, (237, 50, 200)),
        ):
            if point is not None:
                x = left_x + int(point[2] * scale_x)
                y = left_y + int(point[1] * scale_y)
                draw.ellipse((x - 6, y - 6, x + 6, y + 6), fill=color, outline="black", width=2)
        ratios = np.asarray([
            np.nan if row.get(f"bin_{index:02d}_ratio") is None else float(row[f"bin_{index:02d}_ratio"])
            for index in range(1, 11)
        ])
        draw_chart(draw, ratios, None, (100, 650, 1320, 910))
        draw.text((520, 940), "Spine part (shaft -> tip)", fill="black", font=font)
        pages.append(page)
        if progress:
            progress("Exporting validation PDF", position + 1, len(selected), f"Spine {row['spine_id']}")

    if not pages:
        page, draw = page_and_draw()
        draw.text((500, 480), "No spines matched this optional PDF category.", fill="black", font=font)
        pages.append(page)
    pages[0].save(path, "PDF", resolution=150.0, save_all=True, append_images=pages[1:])


def export_measurements(
    manifest: dict[str, object],
    workbook_path: str | Path,
    *,
    validation_pdf: bool = False,
    excluded_audit_pdf: bool = False,
    invalid_audit_pdf: bool = False,
    pdf_margin_um: float = 1.0,
    progress: Progress | None = None,
) -> dict[str, object]:
    path = Path(workbook_path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    tables = collect_export_tables(manifest)
    _write_workbook(path, tables)
    csv_directory = path.parent / f"{path.stem}_csv"
    csv_paths = _write_csvs(csv_directory, tables)
    written_pdfs: list[Path] = []
    indexed_rows: list[tuple[int, dict[str, object]]] = []
    for specimen_index, _specimen in enumerate(manifest["specimens"]):
        try:
            result, _audit = filtered_measurement_result(manifest, specimen_index)
        except ValueError:
            continue
        indexed_rows.extend(
            (specimen_index, row)
            for row in result.get("distribution_rows", [])
            if not bool(row.get("volume_filter_excluded", False))
        )
    selections = []
    if validation_pdf:
        selections.append((path.with_name(f"{path.stem}_distribution_validation.pdf"), [item for item in indexed_rows if bool(item[1].get("spine_valid", True)) and bool(item[1].get("distribution_included", False))], True))
    if excluded_audit_pdf:
        selections.append((path.with_name(f"{path.stem}_distribution_excluded_audit.pdf"), [item for item in indexed_rows if bool(item[1].get("spine_valid", True)) and not bool(item[1].get("distribution_included", False))], False))
    if invalid_audit_pdf:
        selections.append((path.with_name(f"{path.stem}_invalid_spines_audit.pdf"), [item for item in indexed_rows if not bool(item[1].get("manual_spine_valid", item[1].get("spine_valid", True)))], False))
    for pdf_path, selected, summaries in selections:
        _pdf_pages(
            manifest,
            pdf_path,
            selected,
            tables["Distribution_Group"],
            pdf_margin_um,
            include_group_summary=summaries,
            progress=progress,
        )
        written_pdfs.append(pdf_path)

    # Verification happens before the UI may offer any future cache deletion.
    load_workbook(path, read_only=True).close()
    if not csv_paths or any(not csv_path.is_file() for csv_path in csv_paths):
        raise OSError("CSV export verification failed.")
    for pdf_path in written_pdfs:
        with pdf_path.open("rb") as stream:
            if stream.read(4) != b"%PDF":
                raise OSError(f"PDF verification failed: {pdf_path.name}")
    return {
        "workbook": str(path),
        "csv_directory": str(csv_directory),
        "pdfs": [str(item) for item in written_pdfs],
        "verified": True,
    }
